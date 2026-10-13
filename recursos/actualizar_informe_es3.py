"""
Script de Actualización del Informe Oficial CINEAPP
Actualiza el diagrama UML (versión 2 con DAO/SQLite), todas las capturas de pantalla de terminal
y la sección de persistencia SQLite y transacciones atómicas en:
- C:\\Users\\lugog\\OneDrive\\Desktop\\Informe_Sumativa2_CINEAPP.docx
- C:\\Users\\lugog\\OneDrive\\Desktop\\Informe_Sumativa2_CINEAPP.pdf
"""

import os
import shutil
import io
import subprocess
from PIL import Image
import docx
from docx.shared import Inches

DOCX_PATH = r"C:\Users\lugog\OneDrive\Desktop\Informe_Sumativa2_CINEAPP.docx"
PDF_PATH = r"C:\Users\lugog\OneDrive\Desktop\Informe_Sumativa2_CINEAPP.pdf"
BACKUP_PATH = r"C:\Users\lugog\OneDrive\Desktop\Informe_Sumativa2_CINEAPP.docx.bak"
CAPTURAS_DIR = r"C:\Users\lugog\OneDrive\Desktop\Capturas_CINEAPP_ES3"
DIAGRAMA_PATH = r"C:\Users\lugog\OneDrive\Desktop\CINEAPP\Diagrama actualizado.png"


def png_to_jpeg_bytes(png_path, quality=98):
    """Convierte una imagen PNG a bytes JPEG con máxima nitidez."""
    with Image.open(png_path) as im:
        rgb_im = im.convert("RGB")
        buf = io.BytesIO()
        rgb_im.save(buf, format="JPEG", quality=quality, subsampling=0)
        return buf.getvalue()


def actualizar_informe():
    print("Iniciando actualización del informe...")

    # 1. Crear copia de respaldo de seguridad
    if os.path.exists(DOCX_PATH) and not os.path.exists(BACKUP_PATH):
        shutil.copy2(DOCX_PATH, BACKUP_PATH)
        print(f"Respaldo creado en: {BACKUP_PATH}")

    doc = docx.Document(DOCX_PATH)

    # 2. Actualizar Título / Subtítulo si corresponde
    if len(doc.paragraphs) > 3:
        p3 = doc.paragraphs[3]
        if "Sumativa N2" in p3.text or "Sumativa N°2" in p3.text:
            p3.text = "Evaluación Sumativa N°2 y N°3: Persistencia DAO (SQLite), Seguridad y Consumo de APIs"

    # 3. Actualizar Explicación del Diagrama UML (P22)
    p22 = doc.paragraphs[22]
    p22.text = (
        "Explicación de cambios: El diagrama evolucionó desde la primera versión incorporando la arquitectura relacional "
        "completa exigida en la Evaluación Sumativa N°3: (1) Creación del paquete desacoplado dao/ con ConexionDB, "
        "PeliculaDAO, VentaDAO y ClienteDAO; (2) Integración con el motor SQLite relacional (cineapp.db), separando "
        "estrictamente la lógica de dominio (model/) del acceso a datos; (3) Consultas 100% parametrizadas con placeholders '?' "
        "para blindar el sistema contra inyección SQL; y (4) Transacciones atómicas unificadas de venta con cabecera y líneas "
        "de detalle (detalles_venta)."
    )

    # 4. Actualizar Imagen del Diagrama de Clases (rId8)
    print("Actualizando Diagrama UML actualizado (rId8)...")
    diag_jpeg = png_to_jpeg_bytes(DIAGRAMA_PATH, quality=98)
    doc.part.rels["rId8"].target_part._blob = diag_jpeg

    # 5. Mapeo de capturas a actualizar en el documento
    mapeo_capturas = {
        "rId10": "05_Polimorfismo_y_CotizacionDolar.png",
        "rId12": "06_Validacion_Edad_Setter.png",
        "rId14": "03_Dulceria_Comprobante_Final.png",
        "rId16": "01_Regla1_EdadInsuficienteError.png",
        "rId18": "02_Acompanante_Adulto_Autorizado.png",
        "rId19": "04_Regla2_SalaLlenaError.png",
        "rId21": "07_RUT_Modulo11_Validacion.png",
        "rId22": "08_CRUD_SQLite_y_Persistencia_DAO.png",
        "rId23": "10_API_Requests_Timeout_y_Contingencia.png",
        "rId24": "11_Control_Errores_Menu_Opcion99_Letras.png",
    }

    for rid, archivo in mapeo_capturas.items():
        ruta_img = os.path.join(CAPTURAS_DIR, archivo)
        if os.path.exists(ruta_img) and rid in doc.part.rels:
            print(f"Actualizando {rid} con {archivo}...")
            jpeg_bytes = png_to_jpeg_bytes(ruta_img, quality=98)
            doc.part.rels[rid].target_part._blob = jpeg_bytes

    # 6. Actualizar Texto de Sección 5.2 (P110, P111, P115)
    doc.paragraphs[110].text = "5.2 Módulo CRUD de Películas y Persistencia Relacional en SQLite con Arquitectura DAO (P02 a P06)"
    doc.paragraphs[111].text = (
        "Para garantizar que los datos persistan de manera confiable y estructurada al reiniciar el programa (P05) y permitir "
        "la administración de la cartelera, se implementó el paquete dao/ desacoplado de las clases de negocio (model/), con el "
        "componente ConexionDB y PeliculaDAO conectados a la base de datos SQLite 'cineapp.db'. La inicialización crea "
        "automáticamente las tablas del esquema relacional DDL si no existen, con claves foráneas habilitadas (PRAGMA foreign_keys = ON;). "
        "Todas las operaciones se ejecutan mediante consultas 100% parametrizadas con tuplas y placeholders '?', impidiendo cualquier "
        "vulnerabilidad de inyección SQL. El sistema permite:"
    )
    doc.paragraphs[115].text = (
        "P05 (Persistencia): Al cerrar el programa con la opción 0 y volver a ejecutar python main.py, los datos persisten "
        "intactos gracias a la consulta directa sobre las tablas SQLite de cineapp.db gestionadas por PeliculaDAO."
    )
    doc.paragraphs[118].text = (
        "Figura 5.2: Evidencia del ciclo completo CRUD en SQLite y comprobación de persistencia de datos (P02 a P06) mediante PeliculaDAO."
    )

    # 7. Actualizar Sección 5.3 Dólar -> 5.4, y 5.4 Menú -> 5.5
    doc.paragraphs[119].text = "5.4 Indicador Oficial del Dólar (mindicador.cl) con Requests, Timeout y Contingencia Defensiva (P16 y P17)"
    doc.paragraphs[120].text = (
        "Para cumplir con los criterios 3.1.1 y 3.1.3 (pruebas P16 y P17), la clase CotizacionDolar encapsula el consumo del servicio externo "
        "conectándose al endpoint https://mindicador.cl/api/dolar mediante la librería requests, con un tiempo máximo de espera explícito "
        "(timeout=5). El valor obtenido en vivo se inyecta en ProductoImportado para cotizar dinámicamente en pesos chilenos los snacks en USD. "
        "Si la red falla, se produce una demora excesiva (requests.exceptions.Timeout) o error de conexión (ConnectionError), el bloque try/except "
        "defensivo atrapa la excepción granular, emite una advertencia de contingencia en consola y adopta el valor de respaldo referencial ($950 CLP), "
        "asegurando la continuidad operacional del sistema sin caídas ni bloqueos."
    )
    doc.paragraphs[123].text = (
        "Figura 5.4: Evidencia de consulta del indicador externo dólar con requests (timeout=5s) y mecanismo de contingencia sin conexión (P16 y P17)."
    )

    doc.paragraphs[124].text = "5.5 Menú Principal, Control de Opciones Inválidas y Validación de Entradas (P01, P18 y P19)"
    doc.paragraphs[127].text = (
        "Figura 5.5: Evidencia del Menú Principal, control de opción no existente 99 (P18) y captura de letras en campo numérico (P19)."
    )
    doc.paragraphs[128].text = "5.6 Matriz de Cobertura y Cumplimiento del Guion de Pruebas (P01 a P19)"

    # 8. Insertar Nueva Subsección 5.3 (Transacción Atómica VentaDAO) antes de P119 (ahora 5.4)
    # Verificamos si ya existe para evitar duplicar
    ya_existe_tx = any("Transacción Atómica de Venta" in p.text for p in doc.paragraphs)
    if not ya_existe_tx:
        print("Insertando nueva subsección 5.3 de Transacción Atómica VentaDAO...")
        p_ref = doc.paragraphs[119]

        p_h = p_ref.insert_paragraph_before("5.3 Transacción Atómica de Venta y Líneas de Detalle en SQLite (VentaDAO) (P12 y P13)", style="Heading 2")
        p_txt = p_ref.insert_paragraph_before(
            "Para dar cumplimiento estricto a los criterios 2.1.2 y 2.1.3 (pruebas P12 y P13), el sistema unifica la compra de entradas de cine y "
            "productos de dulcería en una única transacción atómica gestionada por VentaDAO. Dentro de un bloque transaccional seguro (with conn:), "
            "se registra primero la cabecera en la tabla 'ventas' y posteriormente cada uno de los ítems adquiridos en la tabla 'detalles_venta' "
            "(vinculados por clave foránea id_venta). Todas las sentencias utilizan consultas parametrizadas con placeholders '?'. En caso de fallo, "
            "SQLite realiza un rollback automático impidiendo registros huérfanos o parciales. Asimismo, el método VentaDAO.obtener_venta_con_detalles() "
            "permite recuperar la venta y reconstruir la totalidad de sus líneas de detalle y subtotales."
        )
        p_img = p_ref.insert_paragraph_before()
        p_img.paragraph_format.space_before = docx.shared.Pt(6)
        p_img.paragraph_format.space_after = docx.shared.Pt(4)
        run_img = p_img.add_run()
        img_tx_path = os.path.join(CAPTURAS_DIR, "09_Transaccion_Unificada_VentaDAO_SQLite.png")
        run_img.add_picture(img_tx_path, width=Inches(6.2))

        p_cap = p_ref.insert_paragraph_before(
            "Figura 5.3: Evidencia de transacción unificada con entradas y dulcería guardada atómicamente en SQLite (VentaDAO) y consulta de líneas de detalle (P12 y P13)."
        )
        p_cap.paragraph_format.space_after = docx.shared.Pt(12)

    # 9. Actualizar Tabla 5 (Filas de Persistencia y Transacciones)
    t5 = doc.tables[5]
    for row in t5.rows:
        cells_text = [c.text.strip() for c in row.cells]
        if len(cells_text) >= 5:
            if cells_text[0] == "P05":
                row.cells[3].text = "El registro sigue guardado intacto en la base de datos SQLite (cineapp.db) vía PeliculaDAO."
            elif cells_text[0] == "P12":
                row.cells[3].text = "Guardado como un solo registro de venta con transacción atómica y detalles en SQLite (VentaDAO)."
            elif cells_text[0] == "P13":
                row.cells[3].text = "Se ven todas sus líneas de detalle y subtotales consultadas desde la base de datos SQLite."

    # 10. Guardar Documento Word
    doc.save(DOCX_PATH)
    print(f"Documento Word actualizado exitosamente: {DOCX_PATH}")


def exportar_a_pdf():
    print("Exportando documento a PDF mediante Word COM...")
    ps_script = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
try {{
    $doc = $word.Documents.Open('{DOCX_PATH}')
    $doc.SaveAs([ref]'{PDF_PATH}', [ref]17)
    $doc.Close()
    Write-Host 'PDF exportado exitosamente.'
}} finally {{
    $word.Quit()
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
    [System.GC]::Collect()
    [System.GC]::WaitForPendingFinalizers()
}}
"""
    ps_file = os.path.join(os.path.dirname(__file__), "export_temp.ps1")
    with open(ps_file, "w", encoding="utf-8") as f:
        f.write(ps_script)

    try:
        proc = subprocess.run(
            ["powershell", "-ExecutionPolicy", "Bypass", "-File", ps_file],
            capture_output=True,
            text=True,
            check=True
        )
        print(proc.stdout)
    except subprocess.CalledProcessError as e:
        print("Error al exportar PDF:")
        print(e.stderr)
        raise
    finally:
        if os.path.exists(ps_file):
            os.remove(ps_file)


if __name__ == "__main__":
    actualizar_informe()
    exportar_a_pdf()
