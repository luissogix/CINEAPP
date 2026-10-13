import docx

doc = docx.Document(r"C:\Users\lugog\OneDrive\Desktop\Informe_Sumativa2_CINEAPP.docx")

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

with open("auditoria_informe.txt", "w", encoding="utf-8") as f:
    f.write("=== PARAGRAPHS ===\n")
    for i, p in enumerate(doc.paragraphs):
        f.write(f"P{i:03d} [{p.style.name}] (len={len(p.text)}): {p.text}\n")
    
    f.write("\n=== TABLES ===\n")
    for t_idx, t in enumerate(doc.tables):
        f.write(f"\nTABLE {t_idx} (rows={len(t.rows)}, cols={len(t.columns)}):\n")
        for r_idx, r in enumerate(t.rows):
            cells = [c.text.strip().replace('\n', ' ') for c in r.cells]
            f.write(f"  R{r_idx:02d}: {cells}\n")

print("Auditoría guardada en auditoria_informe.txt")
